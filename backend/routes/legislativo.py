"""Copiloto legislativo: minutas de PL, indicação, requerimento, moção e justificativa de emenda."""
import io
import json
import re

from flask import Blueprint, g, jsonify, request, send_file

import planos
from auth import login_requerido
from extensions import ErroAPI, db
from models import TIPOS_MINUTA, Minuta
from routes import dados, gabinete_da_conta
from services import ia, kb, legislativo_fontes, prompts, rag_piloto

bp = Blueprint("legislativo", __name__, url_prefix="/api")


def _minuta(mid):
    m = db.session.get(Minuta, mid)
    if not m:
        raise ErroAPI("Minuta não encontrada.", 404)
    gabinete_da_conta(m.gabinete_id)
    return m


@bp.get("/gabinetes/<int:gid>/minutas")
@login_requerido
def listar(gid):
    gab = gabinete_da_conta(gid)
    return jsonify([m.dict(completo=False) for m in Minuta.query.filter_by(gabinete_id=gab.id).order_by(Minuta.id.desc()).all()])


@bp.post("/gabinetes/<int:gid>/minutas")
@login_requerido
def gerar(gid):
    gab = gabinete_da_conta(gid)
    d = dados()
    tipo, demanda = d.get("tipo"), (d.get("demanda") or "").strip()
    if any(k in d for k in ("tema", "objetivo", "publico")):
        if any(not isinstance(d.get(k), str) or not d[k].strip() for k in ("tema", "objetivo", "publico")):
            raise ErroAPI("Preencha tema, objetivo principal e público impactado.")
        demanda = f"Tema: {d['tema'][:500]}\nObjetivo: {d['objetivo'][:2500]}\nPúblico impactado: {d['publico'][:700]}"
    if tipo not in TIPOS_MINUTA:
        raise ErroAPI("Escolha o tipo de proposição.")
    if len(demanda) < 20:
        raise ErroAPI("Descreva a demanda com um pouco mais de detalhe (o problema, quem é atingido, onde).")
    planos.exigir("minutas")
    triagem = rag_piloto.avaliar(gab, tipo, demanda)
    if triagem['estado'] == 'bloqueado':
        return jsonify({'erro': 'Possível impedimento de competência ou iniciativa. Submeta a demanda à assessoria jurídica; avalie indicação ou requerimento.', 'triagem': triagem}), 422
    base = json.dumps(triagem['base_constitucional'], ensure_ascii=False)
    regimento = json.dumps(triagem['fontes'], ensure_ascii=False)
    usuario = prompts.usuario_minuta(tipo, demanda[:4000], gab, base, regimento)
    usuario += "\nFontes ausentes: " + json.dumps(triagem['fontes_ausentes'])
    parecidas = []
    if planos.tem_modulo(g.conta, "legislativo"):
        palavras = re.findall(r"[A-Za-zÀ-ú]{5,}", demanda)[:3]
        parecidas = legislativo_fontes.leis_parecidas([" ".join(palavras)], gab.uf, limite=6) if palavras else []
        if parecidas:
            usuario += ("\n\n=== LEIS PARECIDAS EM OUTROS ENTES (use para fundamentar e comparar; não invente outras) ===\n"
                        + json.dumps(parecidas, ensure_ascii=False)[:8000])
    res, resp = ia.gerar(g.conta.id, "minuta", "redacao", prompts.SISTEMA_MINUTA, usuario, prompts.demo_minuta(tipo, demanda))
    verif = ia.verificar(g.conta.id, "minuta", prompts.SISTEMA_VERIF_MINUTA, prompts.usuario_verif_minuta(usuario, res),
                         resp, prompts.DEMO_VERIF)
    m = Minuta(gabinete_id=gab.id, tipo=tipo, demanda=demanda, titulo=(res.get("titulo") or TIPOS_MINUTA[tipo])[:400],
               texto=res.get("texto") or "", justificativa=res.get("justificativa") or "",
               analise=json.dumps({**(res.get("analise") or {}), "ementa": res.get("ementa"),
                                   "regimento_usado": bool(triagem["fontes"]), "leis_parecidas": parecidas, "triagem": triagem}, ensure_ascii=False),
               verificacao=json.dumps(verif, ensure_ascii=False), modelo=resp.modelo)
    db.session.add(m)
    db.session.commit()
    return jsonify(m.dict()), 201


@bp.get("/minutas/<int:mid>")
@login_requerido
def ver(mid):
    return jsonify(_minuta(mid).dict())


@bp.put("/minutas/<int:mid>")
@login_requerido
def editar(mid):
    m = _minuta(mid)
    d = dados()
    for c in ("titulo", "texto", "justificativa"):
        if c in d:
            setattr(m, c, d[c])
    db.session.commit()
    return jsonify(m.dict())


@bp.delete("/minutas/<int:mid>")
@login_requerido
def excluir(mid):
    m = _minuta(mid)
    from models_piloto import RevisaoMinuta
    RevisaoMinuta.query.filter_by(minuta_id=mid).delete()
    db.session.delete(m)
    db.session.commit()
    return jsonify({"ok": True})


@bp.get("/minutas/<int:mid>/docx")
@login_requerido
def baixar_docx(mid):
    from docx import Document
    from docx.shared import Pt
    m = _minuta(mid)
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name, estilo.font.size = "Arial", Pt(12)
    gab = gabinete_da_conta(m.gabinete_id)
    revisao = rag_piloto.revisao_atual(m, gab)
    doc.add_paragraph("MINUTA — REVISÃO HUMANA PENDENTE" if revisao['estado'] != 'aprovado' else "MINUTA — REVISADA PELA EQUIPE")
    doc.add_heading(m.titulo or TIPOS_MINUTA.get(m.tipo, "Proposição"), level=1)
    for linha in (m.texto or "").split("\n"):
        doc.add_paragraph(linha)
    if m.justificativa:
        doc.add_heading("Justificativa", level=2)
        for linha in m.justificativa.split("\n"):
            doc.add_paragraph(linha)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"minuta-{m.id}.docx",
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


# ------------------------------------------------------------------ Copiloto Legislativo: pesquisa e análise
from models import AnaliseProposicao  # noqa: E402


@bp.get("/gabinetes/<int:gid>/legislativo/pesquisa")
@login_requerido
def pesquisar(gid):
    gab = gabinete_da_conta(gid)
    planos.exigir_modulo("legislativo")
    q = (request.args.get("q") or "").strip()
    fonte = request.args.get("fonte") or "camara"
    if len(q) < 3:
        raise ErroAPI("Digite ao menos 3 letras para pesquisar.")
    tipo = request.args.get("tipo") or "PL"
    if fonte == "camara":
        return jsonify(legislativo_fontes.camara_buscar(q, tipo, request.args.get("ano")))
    if fonte == "senado":
        return jsonify(legislativo_fontes.senado_buscar(q, tipo))
    if fonte == "sapl":
        return jsonify(legislativo_fontes.sapl_buscar(gab.sapl_url, q, request.args.get("ano")))
    if fonte == "leis":
        return jsonify(legislativo_fontes.leis_parecidas([q], gab.uf, limite=15))
    raise ErroAPI("Fonte inválida.")


@bp.get("/gabinetes/<int:gid>/analises")
@login_requerido
def listar_analises(gid):
    gab = gabinete_da_conta(gid)
    itens = AnaliseProposicao.query.filter_by(gabinete_id=gab.id).order_by(AnaliseProposicao.id.desc()).limit(200).all()
    return jsonify([a.dict(completo=False) for a in itens])


@bp.post("/gabinetes/<int:gid>/analises")
@login_requerido
def analisar(gid):
    """Analisa uma proposição de uma base pública (fonte + id_externo) ou um texto colado/enviado pelo gabinete."""
    gab = gabinete_da_conta(gid)
    planos.exigir_modulo("legislativo")
    planos.exigir("analises")
    d = dados()
    fonte = d.get("fonte") or "texto"
    ident, ementa, situacao, url, texto = d.get("identificacao") or "", d.get("ementa") or "", None, d.get("url"), ""
    if fonte == "camara":
        det = legislativo_fontes.camara_detalhe(d.get("id_externo"))
        ident, ementa, situacao, url = det["identificacao"], det["ementa"], det["situacao"], det["url"]
        texto = legislativo_fontes.baixar_texto(det["url_texto"])
    elif fonte == "senado":
        det = legislativo_fontes.senado_detalhe(d.get("id_externo"))
        url = det["url"]
        texto = legislativo_fontes.baixar_texto(det["url_texto"])
    elif fonte == "sapl":
        texto = legislativo_fontes.baixar_texto(d.get("url_texto"))
    else:
        texto = (d.get("texto") or "").strip()
        if len(texto) < 200:
            raise ErroAPI("Cole o texto da proposição (ou envie o PDF) para analisar.")
        ident = ident or "Texto enviado pelo gabinete"
    palavras = [w for w in re.findall(r"[A-Za-zÀ-ú]{5,}", ementa or texto[:600])][:6]
    parecidas = legislativo_fontes.leis_parecidas([" ".join(palavras[:3])] if palavras else [], gab.uf)
    base = kb.base_fixa("01", "02", "03")
    usuario = prompts.usuario_analise(gab, ident, ementa, situacao, texto, parecidas, base)
    res, resp = ia.gerar(g.conta.id, "analise", "analise", prompts.SISTEMA_ANALISE, usuario,
                         prompts.demo_analise(ident, ementa), max_tokens=6000)
    verif = ia.verificar(g.conta.id, "analise", prompts.SISTEMA_VERIF_ANALISE,
                         usuario[:60000] + "\n\n=== ANÁLISE ===\n" + json.dumps(res, ensure_ascii=False), resp, prompts.DEMO_VERIF)
    a = AnaliseProposicao(gabinete_id=gab.id, fonte=fonte, id_externo=d.get("id_externo"), identificacao=ident[:120],
                          ementa=ementa, url=url, situacao=situacao, resultado=json.dumps({**res, "texto_integral": bool(texto)}, ensure_ascii=False),
                          verificacao=json.dumps(verif, ensure_ascii=False), comparacoes=json.dumps(parecidas, ensure_ascii=False),
                          modelo=resp.modelo)
    db.session.add(a)
    db.session.commit()
    return jsonify(a.dict()), 201


@bp.get("/analises/<int:aid>")
@login_requerido
def ver_analise(aid):
    a = db.session.get(AnaliseProposicao, aid)
    if not a:
        raise ErroAPI("Análise não encontrada.", 404)
    gabinete_da_conta(a.gabinete_id)
    return jsonify(a.dict())


@bp.delete("/analises/<int:aid>")
@login_requerido
def excluir_analise(aid):
    a = db.session.get(AnaliseProposicao, aid)
    if not a:
        raise ErroAPI("Análise não encontrada.", 404)
    gabinete_da_conta(a.gabinete_id)
    db.session.delete(a)
    db.session.commit()
    return jsonify({"ok": True})
